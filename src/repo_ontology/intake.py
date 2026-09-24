"""Intake functional requirements from Excel, CSV, or Markdown and translate to ontology specs."""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from rich.console import Console
from rich.panel import Panel
from rich.table import Table

from repo_ontology.loader import OntologyRegistry, load_ontology
from repo_ontology.models import ViewSpec
from repo_ontology.translate import match_views, normalize_text

console = Console()
error_console = Console(stderr=True)


class RequirementItem:
    def __init__(
        self,
        source_sheet: str,
        row_num: int,
        role: Optional[str] = None,
        menu_path: str = "",
        page_code: Optional[str] = None,
        title: str = "",
        description: str = "",
        raw_data: Optional[Dict[str, Any]] = None,
    ):
        self.source_sheet = source_sheet
        self.row_num = row_num
        self.role = role
        self.menu_path = menu_path
        self.page_code = page_code
        self.title = title
        self.description = description
        self.raw_data = raw_data or {}
        self.matched_view: Optional[ViewSpec] = None
        self.match_score: float = 0.0
        self.match_reason: str = ""


def parse_excel_file(path: Path) -> List[RequirementItem]:
    """Parse Excel sheets into RequirementItems by auto-detecting column layouts."""
    try:
        import openpyxl
    except ImportError:
        raise ImportError("openpyxl is required to parse Excel files. Run 'pip install openpyxl' or 'uv add openpyxl'.")

    wb = openpyxl.load_workbook(path, data_only=True)
    items = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        if ws.max_row < 2:
            continue

        # Find header row (search in first 5 rows)
        header_row_idx = None
        headers = []
        for r in range(1, min(6, ws.max_row + 1)):
            row_vals = [str(ws.cell(r, c).value or "").strip() for c in range(1, ws.max_column + 1)]
            # Check if this row looks like a header
            kw_count = sum(1 for v in row_vals if any(k in v for k in ["구분", "분류", "메뉴", "페이지", "기능", "설명", "화면", "요구사항"]))
            if kw_count >= 2:
                header_row_idx = r
                headers = row_vals
                break

        if not header_row_idx:
            # Fallback: assume row 1
            header_row_idx = 1
            headers = [str(ws.cell(1, c).value or "").strip() for c in range(1, ws.max_column + 1)]

        # Map column positions
        col_role = -1
        col_main = -1
        col_sub = -1
        col_minor = -1
        col_page_code = -1
        col_page_name = -1
        col_desc = -1

        for idx, h in enumerate(headers):
            h_clean = h.replace(" ", "")
            if "구분" in h_clean or "역할" in h_clean:
                col_role = idx + 1
            elif "대분류" in h_clean or ("메뉴" in h_clean and col_main == -1):
                col_main = idx + 1
            elif "중분류" in h_clean or ("분류" in h_clean and col_sub == -1):
                col_sub = idx + 1
            elif "소분류" in h_clean:
                col_minor = idx + 1
            elif "페이지번호" in h_clean or "페이지코드" in h_clean or "화면ID" in h_clean:
                col_page_code = idx + 1
            elif "페이지명" in h_clean or "화면명" in h_clean or "기능명" in h_clean:
                col_page_name = idx + 1
            elif "기능설명" in h_clean or "설명" in h_clean or "내용" in h_clean or "요구사항" in h_clean:
                col_desc = idx + 1

        # Iterate rows
        for r in range(header_row_idx + 1, ws.max_row + 1):
            role_val = ws.cell(r, col_role).value if col_role > 0 else None
            main_val = ws.cell(r, col_main).value if col_main > 0 else None
            sub_val = ws.cell(r, col_sub).value if col_sub > 0 else None
            minor_val = ws.cell(r, col_minor).value if col_minor > 0 else None
            page_code_val = ws.cell(r, col_page_code).value if col_page_code > 0 else None
            page_name_val = ws.cell(r, col_page_name).value if col_page_name > 0 else None
            desc_val = ws.cell(r, col_desc).value if col_desc > 0 else None

            # Skip empty rows
            all_vals = [role_val, main_val, sub_val, minor_val, page_code_val, page_name_val, desc_val]
            if not any(v is not None and str(v).strip() for v in all_vals):
                continue

            # Build menu path
            menu_parts = [str(p).strip() for p in [main_val, sub_val, minor_val, page_name_val] if p and str(p).strip() and str(p).strip() != "-"]
            menu_path = " > ".join(menu_parts)

            title = str(page_name_val or minor_val or sub_val or main_val or f"Row {r}").strip()
            desc = str(desc_val or "").strip()

            item = RequirementItem(
                source_sheet=sheet_name,
                row_num=r,
                role=str(role_val).strip() if role_val else None,
                menu_path=menu_path,
                page_code=str(page_code_val).strip() if page_code_val else None,
                title=title,
                description=desc,
                raw_data={"row": r, "headers": headers},
            )
            items.append(item)

    return items


def analyze_requirements(registry: OntologyRegistry, items: List[RequirementItem]) -> Tuple[List[RequirementItem], List[RequirementItem]]:
    """Match items against ontology views and partition into matched vs unmatched."""
    matched = []
    unmatched = []

    for item in items:
        # Search strategy:
        # 1. By page_code if available
        # 2. By menu_path
        # 3. By title
        candidates = []
        if item.page_code:
            candidates.extend(match_views(registry, item.page_code))
        if item.menu_path:
            candidates.extend(match_views(registry, item.menu_path))
        if item.title:
            candidates.extend(match_views(registry, item.title))

        if candidates:
            # Sort and pick top candidate with score >= 50
            candidates.sort(key=lambda x: x[1], reverse=True)
            top_view, top_score, top_reason = candidates[0]
            if top_score >= 50.0:
                item.matched_view = top_view
                item.match_score = top_score
                item.match_reason = top_reason
                matched.append(item)
                continue

        unmatched.append(item)

    return matched, unmatched


def generate_markdown_report(registry: OntologyRegistry, matched: List[RequirementItem], unmatched: List[RequirementItem], source_name: str) -> str:
    """Generate detailed markdown impact report from intake analysis."""
    lines = [
        f"# Requirements Intake & Ontology Translation Report",
        f"",
        f"> **Source**: `{source_name}`",
        f"> **Analysis Result**: Total {len(matched) + len(unmatched)} requirements — **{len(matched)} Matched (Modifications)**, **{len(unmatched)} Unmatched (New Features/Pages)**",
        f"",
        f"---",
        f"",
        f"## 1. Summary Overview",
        f"",
        f"| Category | Count | Description |",
        f"|---|---:|---|",
        f"| **Matched (Existing Views)** | {len(matched)} | Existing views/routes in `.ontology/sitemap.yml`. Ready for impact tracing and TDD. |",
        f"| **Unmatched (New Specs)** | {len(unmatched)} | New features, screens, or menus requiring FSD slice & backend scaffolding. |",
        f"",
        f"---",
        f"",
        f"## 2. Matched Requirements & Impacted Code Paths",
        f"",
    ]

    # Group matched by view
    by_view: Dict[str, List[RequirementItem]] = {}
    for item in matched:
        view_name = item.matched_view.view if item.matched_view else "Unknown"
        by_view.setdefault(view_name, []).append(item)

    for view_name, reqs in by_view.items():
        view = reqs[0].matched_view
        lines.append(f"### View: `{view_name}`" + (f" (`{view.page_code}`)" if view.page_code else ""))
        lines.append(f"- **Menu Path**: `{view.menu}`")
        if view.roles:
            lines.append(f"- **Roles**: {', '.join(view.roles)}")
        
        b = view.binding
        if b:
            if b.frontend:
                lines.append(f"- **Frontend Route**: `{b.frontend.route or '-'}`")
                lines.append(f"- **FSD Slice**: `{b.frontend.slice or '-'}`")
            if b.backend:
                lines.append(f"- **Backend Router**: `{b.backend.router or '-'}`")
                if b.backend.endpoints:
                    lines.append(f"- **API Endpoints**: " + ", ".join(f"`{ep}`" for ep in b.backend.endpoints))
            if b.ontology:
                if b.ontology.objects:
                    lines.append(f"- **Ontology Objects**: " + ", ".join(f"`{o}`" for o in b.ontology.objects))
                if b.ontology.actions:
                    lines.append(f"- **Ontology Actions**: " + ", ".join(f"`{a}`" for a in b.ontology.actions))
                if b.ontology.functions:
                    lines.append(f"- **Ontology Functions**: " + ", ".join(f"`{fn}`" for fn in b.ontology.functions))
                if b.ontology.rules:
                    lines.append(f"- **Bound Rules (Invariants)**: " + ", ".join(f"`{r}`" for r in b.ontology.rules))

        lines.append("")
        lines.append(f"**Associated Requirement Rows ({len(reqs)})**:")
        for r in reqs:
            summary = r.description.split("\n")[0][:100] if r.description else r.title
            lines.append(f"- `[{r.source_sheet} R{r.row_num}]` **{r.title}**: {summary}")
        lines.append("")

    lines.append(f"---")
    lines.append(f"## 3. Unmatched Requirements (Scaffolding Candidates)")
    lines.append(f"")
    lines.append(f"These items do not currently exist in `.ontology/sitemap.yml`. They represent new screens, apps, or modules to scaffold:")
    lines.append(f"")

    for u in unmatched[:30]:  # Cap display to avoid overflow
        summary = u.description.split("\n")[0][:120] if u.description else "-"
        p_code = f"`{u.page_code}` " if u.page_code else ""
        lines.append(f"- `[{u.source_sheet} R{u.row_num}]` {p_code}**{u.menu_path or u.title}** ({u.role or 'All'}): {summary}")

    if len(unmatched) > 30:
        lines.append(f"- *... and {len(unmatched) - 30} more items.*")

    lines.append(f"")
    lines.append(f"---")
    lines.append(f"## 4. Recommended Action Plan")
    lines.append(f"1. **For Matched Items**: Run `otlg trace <Action/Object>` to verify test coverage and execute `/tdd-plan`.")
    lines.append(f"2. **For Unmatched Items**: Create new `ViewSpec` entries in `.ontology/sitemap.yml` and scaffold FSD slices / API routers.")

    return "\n".join(lines)
